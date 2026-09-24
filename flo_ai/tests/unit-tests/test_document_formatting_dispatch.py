"""BaseLLM.format_document_in_message routes documents by mime type.

Office and plain-text formats are extracted to text; PDF (and an unknown mime,
assumed PDF) keeps each provider's native path. These pin that split, the
per-provider text shape, and the shared cache.
"""

import base64
import io
import threading
from unittest.mock import patch

import pytest

from flo_ai import DocumentMessageContent
from flo_ai.llm.anthropic_llm import Anthropic
from flo_ai.llm.gemini_llm import Gemini
from flo_ai.llm.openai_llm import OpenAI
from flo_ai.utils.document_text_extraction import (
    DOC_MIME_TYPE,
    EXTRACTABLE_DOCUMENT_MIME_TYPES,
    DocumentExtractionError,
)

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PAGES = [{'type': 'image_url', 'image_url': {'url': 'page-1'}}]


def _document(data: bytes, mime_type, file_name='f') -> DocumentMessageContent:
    return DocumentMessageContent(
        base64=base64.b64encode(data).decode(),
        mime_type=mime_type,
        file_name=file_name,
    )


def _xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.append(['Region', 'Revenue'])
    workbook.active.append(['East', 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph('Board minutes')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestExtractablePath:
    @pytest.mark.parametrize(
        'data, mime_type, expected',
        [
            (b'region,revenue\nEast,1200\n', 'text/csv', 'East,1200'),
            (b'plain notes', 'text/plain', 'plain notes'),
            (_xlsx(), XLSX, 'East,1200'),
            (_docx(), DOCX, 'Board minutes'),
        ],
    )
    async def test_openai_gets_a_text_block(self, data, mime_type, expected):
        result = await OpenAI(api_key='k').format_document_in_message(
            _document(data, mime_type, 'report')
        )

        assert result[0]['type'] == 'text'
        assert expected in result[0]['text']

    async def test_anthropic_gets_a_text_block_not_a_document_block(self):
        result = await Anthropic(api_key='k').format_document_in_message(
            _document(b'a,b\n1,2\n', 'text/csv')
        )

        assert result == [{'type': 'text', 'text': result[0]['text']}]
        assert '1,2' in result[0]['text']

    async def test_gemini_gets_a_text_part(self):
        """Gemini overrides format_text_content: it takes a Part, not a dict."""
        result = await Gemini(api_key='k').format_document_in_message(
            _document(b'a,b\n1,2\n', 'text/csv')
        )

        assert result.text is not None
        assert '1,2' in result.text

    async def test_file_name_is_carried_in_the_text(self):
        result = await OpenAI(api_key='k').format_document_in_message(
            _document(b'x', 'text/plain', 'quarterly.txt')
        )

        assert 'quarterly.txt' in result[0]['text']

    async def test_mime_is_normalized_before_routing(self):
        """Casing and charset parameters must not flip the routing decision."""
        with patch.object(OpenAI, '_rasterize_pdf_to_images') as rasterize:
            result = await OpenAI(api_key='k').format_document_in_message(
                _document(b'a,b\n1,2\n', 'TEXT/CSV; charset=utf-8')
            )

        rasterize.assert_not_called()
        assert result[0]['type'] == 'text'


class TestNativePathIsUnchanged:
    async def test_pdf_is_rasterized_not_extracted_on_openai(self):
        with patch.object(OpenAI, '_rasterize_pdf_to_images', return_value=PAGES):
            result = await OpenAI(api_key='k').format_document_in_message(
                _document(b'%PDF-1.4', 'application/pdf', file_name=None)
            )

        assert result == PAGES

    async def test_pdf_is_a_native_document_block_on_anthropic(self):
        result = await Anthropic(api_key='k').format_document_in_message(
            _document(b'%PDF-1.4', 'application/pdf')
        )

        assert result[0]['type'] == 'document'

    async def test_missing_mime_takes_the_native_path(self):
        """An unknown mime is assumed PDF; it must never reach a parser."""
        with patch.object(OpenAI, '_rasterize_pdf_to_images', return_value=PAGES):
            result = await OpenAI(api_key='k').format_document_in_message(
                _document(b'%PDF-1.4', None, file_name=None)
            )

        assert result == PAGES

    @pytest.mark.skipif(
        DOC_MIME_TYPE in EXTRACTABLE_DOCUMENT_MIME_TYPES,
        reason='.doc is enabled; this pins the deferred routing only',
    )
    async def test_doc_is_not_extracted_while_deferred(self):
        """.doc is outside the extractable set until its reader lands."""
        with patch.object(OpenAI, '_rasterize_pdf_to_images', return_value=PAGES):
            result = await OpenAI(api_key='k').format_document_in_message(
                _document(b'\xd0\xcf\x11\xe0', 'application/msword', file_name=None)
            )

        assert result == PAGES


class TestCachingAndErrors:
    async def test_extracted_result_is_cached_per_llm_class(self):
        document = _document(b'a,b\n1,2\n', 'text/csv')
        llm = OpenAI(api_key='k')

        with patch(
            'flo_ai.llm.base_llm.extract_document_text', return_value='cached text'
        ) as extract:
            first = await llm.format_document_in_message(document)
            second = await llm.format_document_in_message(document)

        extract.assert_called_once()
        assert first == second
        assert document._formatted_cache['OpenAI'] == first

    async def test_extraction_error_propagates_unwrapped(self):
        """Callers need to tell an unreadable upload from a provider failure."""
        with pytest.raises(DocumentExtractionError):
            await OpenAI(api_key='k').format_document_in_message(
                _document(b'not a zip', DOCX, 'broken.docx')
            )

    @pytest.mark.parametrize(
        'bad_base64', ['!!!!', 'SGVs bG8=', 'data:x;base64,SGVsbG8=']
    )
    async def test_malformed_base64_is_rejected_not_read_as_empty(self, bad_base64):
        """A lenient decode turns `!!!!` into b'' and the model is told the
        file is empty. The upload is broken and should fail as such."""
        document = DocumentMessageContent(
            base64=bad_base64, mime_type='text/plain', file_name='notes.txt'
        )

        with pytest.raises(ValueError, match='not valid base64'):
            await OpenAI(api_key='k').format_document_in_message(document)

    async def test_extraction_runs_off_the_event_loop_thread(self):
        """Parsing a DOCX/XLSX must not stall the caller's event loop.

        This guarantee used to live in wavefront, which parsed documents before
        handing them over. Extraction now happens here, so it is pinned here.
        """
        loop_thread = threading.get_ident()
        extraction_threads = []

        def record_thread(*args, **kwargs):
            extraction_threads.append(threading.get_ident())
            return 'extracted'

        with patch(
            'flo_ai.llm.base_llm.extract_document_text', side_effect=record_thread
        ):
            await OpenAI(api_key='k').format_document_in_message(
                _document(b'a,b\n1,2\n', 'text/csv')
            )

        assert extraction_threads and extraction_threads[0] != loop_thread
